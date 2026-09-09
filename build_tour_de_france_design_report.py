#!/usr/bin/env python3
"""Build an editable DOCX design blueprint for the BT5110 Tour de France project.

The document deliberately separates specification facts, design assumptions, proposed
implementation choices, and items that need stakeholder confirmation.  Diagrams are
generated from code and embedded in Word; all narrative, tables, SQL, and captions
remain ordinary editable DOCX content.

Run from this directory after installing the dependencies listed in requirements.txt:

    python3 build_tour_de_france_design_report.py
"""

from __future__ import annotations

import math
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Circle, FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT_FILE = ROOT / "Tour_de_France_Database_Design_Blueprint.docx"
ASSET_DIR = ROOT / "generated_assets"
NUS_LOGO_FILE = ROOT / "assets" / "nus_logo.png"

SOURCE_FILES = [
    ROOT / "BT5110 Project Tour de France 2026-1.pdf",
    ROOT / "Lectures" / "L1. Introduction.pdf",
    ROOT / "Lectures" / "L2. Relational Model.pdf",
    ROOT / "Lectures" / "L3. Entity-Relationship Model.pdf",
    ROOT / "Lectures" / "L4. FD v2.pdf",
    ROOT / "Tutorials" / "BT5110_Tut2_ER.pdf",
    ROOT / "Tutorials" / "BT5110_Tut3_FD.pdf",
    ROOT / "Tutorials" / "BT5110_Tut4_Normalization.pdf",
    NUS_LOGO_FILE,
]

# Restrained cycling-inspired palette.  The yellow is reserved for decisions and
# warnings so it retains meaning rather than becoming decoration.
INK = "17222B"
NAVY = "12304A"
BLUE = "1D5D7A"
TEAL = "168A83"
YELLOW = "F3C543"
RED = "B44A3C"
GREEN = "2E7D61"
PALE_BLUE = "EAF3F7"
PALE_TEAL = "EAF6F4"
PALE_YELLOW = "FFF7DA"
PALE_RED = "FAECE8"
LIGHT_GREY = "F3F5F6"
MID_GREY = "697780"
DARK_GREY = "3D4A52"
WHITE = "FFFFFF"


def validate_sources() -> None:
    missing = [str(path) for path in SOURCE_FILES if not path.exists()]
    if missing:
        raise FileNotFoundError("Required project/course files are missing:\n" + "\n".join(missing))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("BT5110  |  TOUR DE FRANCE DATABASE  |  DESIGN BLUEPRINT")
    r.font.name = "Liberation Sans"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(MID_GREY)

    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("DESIGN WORKING PAPER   •   ")
    r.font.name = "Liberation Sans"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def add_bottom_border(paragraph, color=TEAL, size="14", space="5") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_update_fields_setting(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    set_repeat_header_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Liberation Sans"

    title = styles["Title"]
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = styles["Subtitle"]
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GREY)

    h1 = styles["Heading 1"]
    h1.font.size = Pt(17)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(NAVY)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(TEAL)
    h3.paragraph_format.space_before = Pt(7)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    for style_name, font_size, color, italic in (
        ("Figure Caption", 8.2, MID_GREY, True),
        ("Small Note", 8.2, MID_GREY, False),
        ("Code Block", 7.4, INK, False),
        ("Source Ref", 8.5, DARK_GREY, False),
    ):
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = "DejaVu Sans Mono" if style_name == "Code Block" else "Liberation Sans"
        style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.italic = italic
        style.paragraph_format.space_after = Pt(3)
        if style_name == "Figure Caption":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style_name == "Code Block":
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            style.paragraph_format.space_after = Pt(0)

    add_update_fields_setting(doc)


def set_portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    set_repeat_header_footer(section)


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    set_repeat_header_footer(section)


def add_page_number_restart(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(22)
    p.add_run().add_picture(str(NUS_LOGO_FILE), width=Inches(2.95))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("BT5110")
    r.font.name = "Liberation Sans"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph("DATA MANAGEMENT AND WAREHOUSING")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.runs[0]
    r.font.name = "Liberation Sans"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(MID_GREY)

    p = doc.add_paragraph("Tour de France\nDatabase & Analytics System", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph("Design strategy, proposed architecture, and validation plan", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)

    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    labels = ["Prepared by", "Group / tutorial", "Academic term", "Document status"]
    values = [
        "[Student names and IDs]",
        "Group 5",
        "AY2026/27 Semester 1",
        "6 September 2026",
    ]
    for i, (label, value) in enumerate(zip(labels, values)):
        left, right = t.rows[i].cells
        left.width = Inches(1.55)
        right.width = Inches(4.35)
        set_cell_margins(left, 70, 100, 70, 100)
        set_cell_margins(right, 70, 100, 70, 100)
        set_cell_shading(left, PALE_BLUE)
        for cell in (left, right):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run(label.upper())
        lr.font.name = "Liberation Sans"
        lr.font.size = Pt(8)
        lr.font.bold = True
        lr.font.color.rgb = RGBColor.from_string(BLUE)
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(0)
        rr = rp.add_run(value)
        rr.font.name = "Liberation Sans"
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SCOPE NOTE")
    r.font.name = "Liberation Sans"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)
    p.add_run("  This paper proposes a defensible design. It does not invent race data, dashboard screenshots, or benchmark timings.")

    p.add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph("Contents", style="Heading 1")
    add_bottom_border(heading)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and update this field if the contents are not shown."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, sep, text, end):
        run._r.append(node)
    p.paragraph_format.space_after = Pt(8)
    note = doc.add_paragraph(
        "The report is organised as a design argument: requirements → conceptual model → relational integrity → analytics model → evidence plan.",
        style="Small Note",
    )
    p = doc.add_paragraph("CONTENTS AT A GLANCE")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.runs[0]
    r.font.name = "Liberation Sans"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    add_table(
        doc,
        ["Part", "Purpose"],
        [
            ["Executive design position", "Recommended two-model architecture and five headline findings"],
            ["1. Scope and requirements", "Traceability, explicit assumptions, and questions to confirm"],
            ["2. Conceptual design", "Core ER model, cardinalities, and classification extension"],
            ["3. Relational design", "Relations, keys, constraints, delete/load policy"],
            ["4. Normalization", "Functional dependencies, anomaly avoidance, BCNF conclusion"],
            ["5. Analytics architecture", "Fact grain, star schema, dimensions, ETL quality gates"],
            ["6. Dashboard design", "Page architecture, OLAP behaviour, interaction standards"],
            ["7. Analytical evidence", "Representative SQL and normalized-vs-star benchmark method"],
            ["8. Validation and delivery", "Test suite, timeline, group workflow, risk register"],
            ["9. Proposed findings", "Design conclusion and boundaries of current evidence"],
            ["Appendices", "DDL starter, multi-edition path, source map, submission checklist"],
        ],
        widths=[2.15, 4.45],
        font_size=7.6,
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        add_bottom_border(p)


def add_text(doc: Document, text: str, *, bold_lead: str | None = None, italic=False) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        first.bold = True
        first.font.color.rgb = RGBColor.from_string(NAVY)
        second = p.add_run(text[len(bold_lead) :])
        second.italic = italic
    else:
        r = p.add_run(text)
        r.italic = italic


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2.5)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(doc: Document, label: str, text: str, kind: str = "decision") -> None:
    colors = {
        "decision": (PALE_YELLOW, YELLOW),
        "finding": (PALE_TEAL, TEAL),
        "risk": (PALE_RED, RED),
        "note": (PALE_BLUE, BLUE),
    }
    fill, accent = colors[kind]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(1.25)
    right.width = Inches(5.45)
    set_cell_shading(left, accent)
    set_cell_shading(right, fill)
    for cell in (left, right):
        set_cell_margins(cell, 100, 110, 100, 110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_after = Pt(0)
    lr = lp.add_run(label.upper())
    lr.font.name = "Liberation Sans"
    lr.font.size = Pt(8)
    lr.font.bold = True
    lr.font.color.rgb = RGBColor.from_string(WHITE if accent != YELLOW else INK)
    rp = right.paragraphs[0]
    rp.paragraph_format.space_after = Pt(0)
    rr = rp.add_run(text)
    rr.font.name = "Liberation Sans"
    rr.font.size = Pt(8.8)
    rr.font.color.rgb = RGBColor.from_string(INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    font_size: float = 8.1,
    alignments: list[int] | None = None,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = widths is None
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for j, label in enumerate(headers):
        cell = header.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 85, 85, 85, 85)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[j])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.bold = True
        r.font.name = "Liberation Sans"
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(WHITE)
    for i, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for j, value in enumerate(values):
            cell = row.cells[j]
            set_cell_margins(cell, 75, 85, 75, 85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cell, LIGHT_GREY)
            if widths:
                cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = alignments[j] if alignments else (WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(value))
            r.font.name = "Liberation Sans"
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption, style="Figure Caption")
    cp.paragraph_format.keep_with_next = False


def add_code_block(doc: Document, code: str, title: str | None = None) -> None:
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Liberation Sans"
        r.font.size = Pt(8.7)
        r.font.color.rgb = RGBColor.from_string(BLUE)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREY)
    set_cell_margins(cell, 105, 125, 105, 125)
    cell.text = ""
    for i, line in enumerate(code.strip("\n").splitlines()):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = "Code Block"
        p.add_run(line if line else " ")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _save_figure(fig, path: Path) -> None:
    fig.savefig(path, dpi=260, bbox_inches="tight", pad_inches=0.08, facecolor="white")
    plt.close(fig)


def draw_design_flow(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.4, 3.35), dpi=240)
    ax.set_xlim(0, 12.4)
    ax.set_ylim(0, 3.35)
    ax.axis("off")
    stages = [
        (0.2, "1", "Read requirements", "Facts, gaps,\nacceptance tests", BLUE),
        (2.65, "2", "Conceptual model", "Entities, keys,\nroles, min/max", TEAL),
        (5.1, "3", "Relational design", "FDs, BCNF,\nconstraints", NAVY),
        (7.55, "4", "Analytics model", "Facts, dimensions,\nETL checks", BLUE),
        (10.0, "5", "Evidence", "Dashboard, SQL,\nbenchmarks", TEAL),
    ]
    for x, n, title, sub, color in stages:
        box = FancyBboxPatch((x, 0.92), 2.15, 1.58, boxstyle="round,pad=0.03,rounding_size=0.08", fc=f"#{WHITE}", ec=f"#{color}", lw=2)
        ax.add_patch(box)
        ax.add_patch(plt.Circle((x + 0.28, 2.23), 0.17, color=f"#{color}"))
        ax.text(x + 0.28, 2.23, n, color="white", ha="center", va="center", fontsize=9, weight="bold")
        ax.text(x + 1.08, 1.88, title, ha="center", va="center", fontsize=10, weight="bold", color=f"#{INK}")
        ax.text(x + 1.08, 1.35, sub, ha="center", va="center", fontsize=8.4, color=f"#{DARK_GREY}")
    for x in [2.37, 4.82, 7.27, 9.72]:
        ax.add_patch(FancyArrowPatch((x, 1.71), (x + 0.22, 1.71), arrowstyle="-|>", mutation_scale=14, lw=1.6, color=f"#{MID_GREY}"))
    ax.annotate("Trace every design choice back to a requirement; feed failed tests back into the model.", xy=(10.98, 0.72), xytext=(6.2, 0.19),
                ha="center", va="center", fontsize=8.5, color=f"#{RED}",
                arrowprops=dict(arrowstyle="-|>", color=f"#{RED}", lw=1.4, connectionstyle="arc3,rad=-0.25"))
    _save_figure(fig, path)


def _entity_box(ax, x, y, w, h, title, lines, color=BLUE, dependent=False, derived=False):
    box = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.012,rounding_size=0.025",
        fc=f"#{WHITE if not derived else PALE_YELLOW}",
        ec=f"#{color}",
        lw=1.8,
        linestyle="--" if dependent else "-",
        zorder=3,
    )
    ax.add_patch(box)
    ax.add_patch(Rectangle((x, y + h - 0.27), w, 0.27, fc=f"#{color}", ec=f"#{color}", zorder=4))
    ax.text(x + 0.04, y + h - 0.135, title, color="white", fontsize=8.2, weight="bold", va="center", zorder=5)
    for i, line in enumerate(lines):
        ax.text(x + 0.05, y + h - 0.36 - i * 0.175, line, fontsize=6.65, color=f"#{INK}", va="top", family="DejaVu Sans Mono", zorder=5)


def _polyline(ax, points, label, label_xy, color=MID_GREY, dash=False):
    xs, ys = zip(*points)
    ax.plot(xs, ys, color=f"#{color}", lw=1.35, ls="--" if dash else "-", zorder=1)
    if label:
        ax.text(label_xy[0], label_xy[1], label, fontsize=6.6, color=f"#{DARK_GREY}", ha="center", va="center",
                bbox=dict(boxstyle="round,pad=0.13", fc="white", ec="none", alpha=0.92), zorder=2)


def _chen_entity(ax, x, y, w, h, label, associative=False):
    """Draw a lecture-style ER entity rectangle centred on (x, y)."""
    ax.add_patch(Rectangle((x - w / 2, y - h / 2), w, h, fc="white", ec="black", lw=1.7, zorder=4))
    if associative:
        ax.add_patch(Rectangle((x - w / 2 + 0.07, y - h / 2 + 0.07), w - 0.14, h - 0.14,
                               fill=False, ec="black", lw=1.0, zorder=4))
    ax.text(x, y, label, ha="center", va="center", fontsize=8.2, color="black", zorder=5)


def _chen_diamond(ax, x, y, w, h, label):
    """Draw a named relationship diamond centred on (x, y)."""
    points = [(x, y + h / 2), (x + w / 2, y), (x, y - h / 2), (x - w / 2, y)]
    ax.add_patch(Polygon(points, closed=True, fc="white", ec="black", lw=1.55, zorder=4))
    ax.text(x, y, label, ha="center", va="center", fontsize=6.3, color="black", zorder=5)


def _chen_edge(ax, points, cardinality, card_xy):
    """Connect an entity and relationship, placing (min, max) near the entity end."""
    xs, ys = zip(*points)
    ax.plot(xs, ys, color="black", lw=1.35, zorder=1)
    ax.text(card_xy[0], card_xy[1], cardinality, ha="center", va="center", fontsize=6.0,
            color="black", bbox=dict(fc="white", ec="none", pad=0.06), zorder=3)


def _chen_attribute(ax, anchor, circle_xy, label, key=False, derived=False, label_xy=None):
    """Draw the lecture's circle attribute notation; dashed means calculated."""
    # Orthogonal attribute stems keep dense groups legible and remain easy to mirror
    # in the editable Draw.io version.
    ax.plot([anchor[0], circle_xy[0], circle_xy[0]], [anchor[1], anchor[1], circle_xy[1]], color="black", lw=1.05,
            ls="--" if derived else "-", zorder=1)
    ax.add_patch(Circle(circle_xy, 0.10, fc="black" if key else "white", ec="black", lw=1.25,
                        ls="--" if derived else "-", zorder=4))
    if label_xy is None:
        label_xy = (circle_xy[0], circle_xy[1] + 0.22)
    ax.text(label_xy[0], label_xy[1], label, ha="center", va="center", fontsize=5.3,
            color="black", zorder=5)


def draw_core_er(path: Path) -> None:
    # Plain, orthogonal Chen/min-max notation matching Lecture 3 and Tutorial 2.
    fig, ax = plt.subplots(figsize=(17.2, 10.2), dpi=260)
    ax.set_xlim(0, 18.8)
    ax.set_ylim(0.25, 11.1)
    ax.axis("off")
    ax.text(0.20, 10.82, "TASK 1 — TOUR DE FRANCE ENTITY-RELATIONSHIP DIAGRAM", fontsize=13,
            weight="bold", color="black")

    # Required entities from the project brief. All use an ordinary single border;
    # the composite identifiers of Result and Exit are shown as key attributes.
    _chen_entity(ax, 2.00, 8.20, 1.70, 0.75, "Country")
    _chen_entity(ax, 6.00, 8.20, 1.70, 0.75, "Team")
    _chen_entity(ax, 10.00, 8.20, 1.70, 0.75, "Rider")
    _chen_entity(ax, 15.90, 8.20, 2.40, 0.82, "Classification / Award")
    _chen_entity(ax, 2.00, 3.00, 1.70, 0.75, "Location")
    _chen_entity(ax, 6.00, 3.00, 1.70, 0.75, "Stage")
    _chen_entity(ax, 10.00, 3.00, 1.70, 0.75, "Result")
    _chen_entity(ax, 15.00, 3.00, 1.70, 0.75, "Exit")

    diamonds = {
        "team_country": (4.00, 8.20, 1.50, 0.70, "based in"),
        "team_rider": (8.00, 8.20, 1.50, 0.70, "has riders"),
        "rider_award": (12.75, 8.20, 1.85, 0.70, "rider award"),
        "team_award": (11.00, 10.05, 1.80, 0.70, "team award"),
        "country_rider": (6.00, 6.50, 1.85, 0.72, "represents country"),
        "country_location": (2.00, 5.55, 1.65, 0.72, "located in"),
        "start": (4.00, 3.95, 1.55, 0.70, "starts at"),
        "finish": (4.00, 2.05, 1.65, 0.70, "finishes at"),
        "rider_result": (10.00, 5.55, 1.55, 0.70, "records"),
        "stage_result": (8.00, 3.00, 1.50, 0.70, "has result"),
        "rider_exit": (13.00, 5.55, 1.60, 0.70, "has exit"),
        "stage_exit": (10.50, 0.78, 1.85, 0.66, "exit before"),
    }
    for values in diamonds.values():
        _chen_diamond(ax, *values)

    # Main top row.
    _chen_edge(ax, [(2.85, 8.20), (3.25, 8.20)], "(0, N)", (3.02, 8.44))
    _chen_edge(ax, [(4.75, 8.20), (5.15, 8.20)], "(1, 1)", (4.98, 8.44))
    _chen_edge(ax, [(6.85, 8.20), (7.25, 8.20)], "(1, N)", (7.02, 8.44))
    _chen_edge(ax, [(8.75, 8.20), (9.15, 8.20)], "(1, 1)", (8.98, 8.44))
    _chen_edge(ax, [(10.85, 8.20), (11.83, 8.20)], "(0, N)", (11.12, 8.44))
    _chen_edge(ax, [(13.68, 8.20), (14.70, 8.20)], "(0, 1)", (14.42, 8.44))

    # Team award is routed above the row using only vertical/horizontal segments.
    _chen_edge(ax, [(6.25, 8.58), (6.25, 10.05), (10.10, 10.05)], "(0, N)", (6.52, 9.68))
    _chen_edge(ax, [(11.90, 10.05), (15.90, 10.05), (15.90, 8.61)], "(0, 1)", (15.62, 9.68))

    # Country-to-Rider nationality is independent of Team nationality.
    _chen_edge(ax, [(2.85, 8.00), (3.20, 8.00), (3.20, 6.50), (5.08, 6.50)],
               "(0, N)", (3.48, 7.72))
    _chen_edge(ax, [(6.92, 6.50), (9.68, 6.50), (9.68, 7.83)],
               "(1, 1)", (9.40, 7.46))

    # Country-to-Location and the two endpoint roles.
    _chen_edge(ax, [(2.00, 7.83), (2.00, 5.91)], "(0, N)", (2.28, 7.45))
    _chen_edge(ax, [(2.00, 5.19), (2.00, 3.38)], "(1, 1)", (2.28, 3.76))
    _chen_edge(ax, [(2.85, 3.20), (3.05, 3.20), (3.05, 3.95), (3.23, 3.95)],
               "(0, N)", (2.98, 3.48))
    _chen_edge(ax, [(4.78, 3.95), (4.95, 3.95), (4.95, 3.20), (5.15, 3.20)],
               "(1, 1)", (5.05, 3.48))
    _chen_edge(ax, [(2.85, 2.80), (3.05, 2.80), (3.05, 2.05), (3.18, 2.05)],
               "(0, N)", (2.98, 2.52))
    _chen_edge(ax, [(4.83, 2.05), (4.95, 2.05), (4.95, 2.80), (5.15, 2.80)],
               "(1, 1)", (5.05, 2.52))

    # Rider/Stage to Result.
    _chen_edge(ax, [(10.00, 7.83), (10.00, 5.90)], "(0, N)", (10.28, 7.43))
    _chen_edge(ax, [(10.00, 5.20), (10.00, 3.38)], "(1, 1)", (10.28, 3.78))
    _chen_edge(ax, [(6.85, 3.00), (7.25, 3.00)], "(0, N)", (7.02, 3.24))
    _chen_edge(ax, [(8.75, 3.00), (9.15, 3.00)], "(1, 1)", (8.98, 3.24))

    # Exit is a required entity associated with both Rider and Stage.
    _chen_edge(ax, [(10.85, 8.00), (11.20, 8.00), (11.20, 5.55), (12.20, 5.55)],
               "(0, 1)", (11.48, 7.68))
    _chen_edge(ax, [(13.80, 5.55), (15.00, 5.55), (15.00, 3.38)],
               "(1, 1)", (14.72, 3.78))
    _chen_edge(ax, [(6.70, 2.63), (7.00, 2.63), (7.00, 0.78), (9.58, 0.78)],
               "(0, N)", (7.28, 2.28))
    _chen_edge(ax, [(11.42, 0.78), (15.00, 0.78), (15.00, 2.63)],
               "(1, 1)", (14.72, 2.24))

    # Entity attributes. Filled circles denote primary-key attributes, as in Lecture 3.
    _chen_attribute(ax, (1.15, 8.35), (0.55, 8.75), "country_code", key=True, label_xy=(0.55, 9.00))
    _chen_attribute(ax, (1.15, 8.05), (0.55, 7.65), "country_name", label_xy=(0.55, 7.38))
    _chen_attribute(ax, (5.55, 8.58), (5.35, 9.18), "team_name", key=True, label_xy=(5.35, 9.43))
    _chen_attribute(ax, (9.45, 8.58), (9.20, 9.18), "bib_no", key=True, label_xy=(9.20, 9.43))
    _chen_attribute(ax, (10.00, 8.58), (10.00, 9.18), "rider_name", label_xy=(10.00, 9.43))
    _chen_attribute(ax, (10.55, 8.58), (10.80, 9.18), "date_of_birth", label_xy=(10.80, 9.43))

    _chen_attribute(ax, (2.00, 2.63), (2.00, 1.88), "location_name", key=True, label_xy=(2.00, 1.61))
    _chen_attribute(ax, (5.45, 2.63), (5.00, 1.86), "stage_id", key=True, label_xy=(5.00, 1.59))
    _chen_attribute(ax, (5.82, 2.63), (5.75, 1.86), "stage_date", label_xy=(5.75, 1.59))
    _chen_attribute(ax, (6.18, 2.63), (6.55, 1.86), "distance_km", label_xy=(6.55, 1.59))
    _chen_attribute(ax, (6.55, 2.63), (7.35, 1.86), "stage_type", label_xy=(7.35, 1.59))

    # Result uses a normal border. Its composite key is shown by the two filled circles.
    result_attrs = [
        (8.20, "stage_id", True), (8.95, "bib_no", True), (9.70, "finish_time", False),
        (10.45, "rank", False), (11.25, "time_bonus", False), (12.20, "time_penalty", False),
    ]
    for i, (x, label, key) in enumerate(result_attrs):
        anchor_x = 9.35 + i * 0.25
        _chen_attribute(ax, (anchor_x, 2.63), (x, 1.82), label, key=key, label_xy=(x, 1.53))
    _chen_attribute(ax, (10.85, 3.00), (11.45, 3.00), "adjusted_time (derived)", derived=True,
                    label_xy=(12.28, 3.00))
    ax.text(10.00, 2.73, "associative entity", ha="center", va="center", fontsize=5.8, zorder=6)

    _chen_attribute(ax, (15.85, 3.28), (16.35, 3.62), "bib_no", key=True, label_xy=(16.82, 3.62))
    _chen_attribute(ax, (15.85, 3.00), (16.35, 3.00), "exit_before_stage_id", label_xy=(17.10, 3.00))
    _chen_attribute(ax, (15.85, 2.72), (16.35, 2.38), "reason", label_xy=(16.78, 2.38))

    _chen_attribute(ax, (15.25, 8.61), (14.85, 9.18), "classification_code", key=True,
                    label_xy=(14.70, 9.43))
    _chen_attribute(ax, (15.70, 8.61), (15.75, 9.18), "name", label_xy=(15.75, 9.43))
    _chen_attribute(ax, (16.15, 8.61), (16.65, 9.18), "subject_type", label_xy=(16.65, 9.43))
    _chen_attribute(ax, (16.60, 8.61), (17.45, 9.18), "measure", label_xy=(17.45, 9.43))
    _chen_attribute(ax, (17.10, 8.20), (17.70, 8.20), "calculation_rule", label_xy=(17.80, 7.93))

    # Cross-relationship/temporal rules that ordinary binary min/max labels cannot show.
    ax.text(2.75, 4.42, "Constraint: every Location is used as\na start or finish at least once.",
            fontsize=6.2, ha="left", va="center")
    ax.text(15.88, 5.18, "Constraint: after an Exit at Stage S,\nthat Rider has no Result for S onward.",
            fontsize=6.2, ha="left", va="center")
    _save_figure(fig, path)


def draw_classification_model(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.8, 5.8), dpi=250)
    ax.set_xlim(0, 12.8)
    ax.set_ylim(0, 5.8)
    ax.axis("off")
    ax.text(0.2, 5.48, "CLASSIFICATION / AWARD EXTENSION", fontsize=14, weight="bold", color=f"#{NAVY}")
    ax.text(12.55, 5.48, "Store evidence; derive standings", fontsize=8.5, color=f"#{MID_GREY}", ha="right")
    _entity_box(ax, 0.4, 2.2, 2.45, 2.0, "CLASSIFICATION", ["PK classification_code", "   name", "   subject_kind", "   measure_kind", "   ranking_direction", "   eligibility_rule"], BLUE)
    _entity_box(ax, 3.7, 2.45, 2.35, 1.65, "SCORING_EVENT", ["PK event_id", "FK stage_id", "FK classification_code", "   event_label", "   route_km"], TEAL)
    _entity_box(ax, 7.0, 2.45, 2.35, 1.65, "RIDER_SCORE", ["PK/FK event_id", "PK/FK bib_no", "   points_awarded"], NAVY, dependent=True)
    _entity_box(ax, 10.05, 3.35, 2.35, 1.35, "RIDER", ["PK bib_no", "   identity + team"], NAVY)
    _entity_box(ax, 10.05, 1.05, 2.35, 1.35, "STAGE / RESULT", ["stage chronology", "adjusted seconds"], TEAL)
    _entity_box(ax, 4.3, 0.2, 3.4, 1.35, "DERIVED STANDINGS / AWARDS", ["GC + Young: result totals", "Points + KOM: score totals", "Best Team: top-three rule"], YELLOW, derived=True)
    _polyline(ax, [(2.85, 3.2), (3.7, 3.2)], "defines\n1   |   0..N", (3.27, 3.55), TEAL)
    _polyline(ax, [(6.05, 3.25), (7.0, 3.25)], "awards\n1   |   0..N", (6.52, 3.58), TEAL)
    _polyline(ax, [(9.35, 3.55), (10.05, 4.02)], "to  N | 1", (9.7, 4.02))
    _polyline(ax, [(6.05, 2.75), (9.7, 2.75), (9.7, 1.75), (10.05, 1.75)], "occurs in  N | 1", (8.15, 2.98))
    _polyline(ax, [(1.65, 2.2), (1.65, 0.88), (4.3, 0.88)], "rule metadata", (2.85, 1.08), YELLOW, True)
    _polyline(ax, [(8.18, 2.45), (8.18, 0.88), (7.7, 0.88)], "raw points", (8.22, 1.55), YELLOW, True)
    _polyline(ax, [(10.05, 1.62), (7.7, 1.15)], "times + eligibility", (8.85, 1.55), YELLOW, True)
    ax.text(0.42, 0.42, "Why this shape?", fontsize=8.3, weight="bold", color=f"#{RED}")
    ax.text(0.42, 0.13, "Points and KOM cannot be reconstructed from finish times alone; GC and Best Team can.", fontsize=7.4, color=f"#{DARK_GREY}")
    _save_figure(fig, path)


def draw_star_schema(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.7, 6.8), dpi=250)
    ax.set_xlim(0, 12.7)
    ax.set_ylim(0, 6.8)
    ax.axis("off")
    ax.text(0.2, 6.48, "ANALYTICAL STAR SCHEMA", fontsize=14, weight="bold", color=f"#{NAVY}")
    ax.text(12.5, 6.48, "Grain: one completed rider × stage", fontsize=8.5, color=f"#{MID_GREY}", ha="right")
    _entity_box(ax, 4.55, 2.15, 3.0, 2.75, "FACT_RESULT", ["PK result_key", "FK rider_key", "FK team_key", "FK stage_key", "FK rider_country_key", "FK team_country_key", "   finish_seconds", "   bonus_seconds", "   penalty_seconds", "   adjusted_seconds", "   finish_rank", "   completed_count (=1)"], NAVY)
    _entity_box(ax, 0.35, 4.55, 2.85, 1.55, "DIM_RIDER", ["PK rider_key", "   bib_no, name, dob", "   age / young flag", "   valid_from / valid_to"], BLUE)
    _entity_box(ax, 0.35, 0.65, 2.85, 1.55, "DIM_TEAM", ["PK team_key", "   team_name", "   leader / role context", "   valid_from / valid_to"], TEAL)
    _entity_box(ax, 9.45, 4.55, 2.85, 1.55, "DIM_STAGE", ["PK stage_key", "   stage_id, date, type", "   start / finish location", "   distance, sequence"], BLUE)
    _entity_box(ax, 9.45, 0.65, 2.85, 1.55, "DIM_COUNTRY  [role-playing]", ["PK country_key", "   ISO code, name", "roles: rider / team", "optionally start / finish"], TEAL)
    for p1, p2, label, xy in [
        ((4.55, 4.35), (3.2, 5.18), "N : 1", (3.85, 4.9)),
        ((4.55, 2.5), (3.2, 1.42), "N : 1", (3.8, 1.95)),
        ((7.55, 4.35), (9.45, 5.18), "N : 1", (8.52, 4.92)),
        ((7.55, 2.5), (9.45, 1.42), "N : 1 (two roles)", (8.62, 1.95)),
    ]:
        _polyline(ax, [p1, p2], label, xy, MID_GREY)
    ax.text(4.72, 1.45, "Optional second fact", fontsize=8.3, weight="bold", color=f"#{RED}")
    ax.text(4.72, 1.11, "FACT_CLASSIFICATION_SCORE at event × rider grain\nfor Points / KOM audits and snapshot analysis.", fontsize=7.4, color=f"#{DARK_GREY}")
    _save_figure(fig, path)


def draw_dashboard_wireframe(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13.2, 7.4), dpi=250)
    ax.set_xlim(0, 13.2)
    ax.set_ylim(0, 7.4)
    ax.axis("off")
    ax.add_patch(Rectangle((0.25, 0.25), 12.7, 6.85, fc="#FBFCFD", ec=f"#{NAVY}", lw=1.6))
    ax.add_patch(Rectangle((0.25, 6.35), 12.7, 0.75, fc=f"#{NAVY}", ec=f"#{NAVY}"))
    ax.text(0.58, 6.73, "TOUR PERFORMANCE COMMAND CENTRE", fontsize=13, weight="bold", color="white", va="center")
    ax.text(12.55, 6.73, "OVERVIEW  |  CLASSIFICATIONS  |  TEAMS & COUNTRIES", fontsize=7.4, color="#DCE8EF", ha="right", va="center")
    ax.add_patch(Rectangle((0.5, 5.75), 11.9, 0.4, fc=f"#{PALE_BLUE}", ec="none"))
    ax.text(0.7, 5.95, "Filters: stage range ▼     stage type ▼     team ▼     rider ▼     nationality ▼", fontsize=8.2, color=f"#{DARK_GREY}", va="center")
    # KPI cards
    cards = [(0.5, "Leader", "Rider A"), (2.75, "GC gap", "+00:41"), (5.0, "Stages", "12 / 21"), (7.25, "Active riders", "154"), (9.5, "Fastest team", "Team X")]
    for x, label, value in cards:
        ax.add_patch(FancyBboxPatch((x, 4.95), 1.95, 0.65, boxstyle="round,pad=0.02,rounding_size=0.04", fc="white", ec="#D7E0E5", lw=1))
        ax.text(x + 0.12, 5.38, label.upper(), fontsize=6.6, color=f"#{MID_GREY}")
        ax.text(x + 0.12, 5.08, value, fontsize=10, weight="bold", color=f"#{NAVY}")
    # Main cumulative line chart
    ax.add_patch(Rectangle((0.5, 2.45), 7.4, 2.25, fc="white", ec="#D7E0E5", lw=1))
    ax.text(0.72, 4.46, "Cumulative adjusted time / gap by stage", fontsize=8.8, weight="bold", color=f"#{INK}")
    xs = [0.95, 1.9, 2.85, 3.8, 4.75, 5.7, 6.65, 7.5]
    for k, color in enumerate([BLUE, TEAL, YELLOW]):
        ys = [3.85 - 0.16 * i + 0.12 * math.sin(i + k) - k * 0.18 for i in range(len(xs))]
        ax.plot(xs, ys, lw=2, color=f"#{color}")
    ax.plot([0.9, 7.55], [2.85, 2.85], color="#DDE4E8", lw=0.8)
    # Ranking panel
    ax.add_patch(Rectangle((8.15, 2.45), 4.25, 2.25, fc="white", ec="#D7E0E5", lw=1))
    ax.text(8.38, 4.46, "Classification comparison", fontsize=8.8, weight="bold", color=f"#{INK}")
    for i, (name, col, pct) in enumerate([("GC", NAVY, 0.82), ("POINTS", TEAL, 0.68), ("KOM", RED, 0.57), ("YOUNG", YELLOW, 0.45)]):
        y = 4.0 - i * 0.36
        ax.text(8.4, y, name, fontsize=6.8, color=f"#{DARK_GREY}", va="center")
        ax.add_patch(Rectangle((9.25, y - 0.08), 2.55, 0.16, fc="#EDF1F3", ec="none"))
        ax.add_patch(Rectangle((9.25, y - 0.08), 2.55 * pct, 0.16, fc=f"#{col}", ec="none"))
    # Bottom panels
    ax.add_patch(Rectangle((0.5, 0.55), 4.45, 1.6, fc="white", ec="#D7E0E5", lw=1))
    ax.text(0.72, 1.91, "Stage profile & rank movement", fontsize=8.6, weight="bold", color=f"#{INK}")
    ax.fill_between([0.85, 1.45, 2.0, 2.55, 3.2, 3.75, 4.55], [0.82, 1.0, 0.92, 1.58, 1.08, 1.3, 0.84], 0.78, color=f"#{PALE_TEAL}")
    ax.plot([0.85, 1.45, 2.0, 2.55, 3.2, 3.75, 4.55], [0.82, 1.0, 0.92, 1.58, 1.08, 1.3, 0.84], color=f"#{TEAL}", lw=1.5)
    ax.add_patch(Rectangle((5.2, 0.55), 3.35, 1.6, fc="white", ec="#D7E0E5", lw=1))
    ax.text(5.42, 1.91, "Country participation", fontsize=8.6, weight="bold", color=f"#{INK}")
    for i, v in enumerate([0.85, 0.63, 0.55, 0.42, 0.34]):
        ax.add_patch(Rectangle((5.55, 1.55 - i * 0.22), 2.3 * v, 0.12, fc=f"#{BLUE}", ec="none"))
    ax.add_patch(Rectangle((8.8, 0.55), 3.6, 1.6, fc=f"#{PALE_YELLOW}", ec="#E8D48A", lw=1))
    ax.text(9.02, 1.91, "Context panel", fontsize=8.6, weight="bold", color=f"#{INK}")
    ax.text(9.02, 1.56, "• selected rider / team\n• exit status and reason\n• data freshness + rule definitions\n• drill-through action", fontsize=7.3, color=f"#{DARK_GREY}", va="top")
    _save_figure(fig, path)


def generate_assets() -> None:
    ASSET_DIR.mkdir(exist_ok=True)
    draw_design_flow(ASSET_DIR / "design_flow.png")
    draw_core_er(ASSET_DIR / "core_er.png")
    draw_classification_model(ASSET_DIR / "classification_model.png")
    draw_star_schema(ASSET_DIR / "star_schema.png")
    draw_dashboard_wireframe(ASSET_DIR / "dashboard_wireframe.png")


def add_executive_summary(doc: Document) -> None:
    add_heading(doc, "Executive design position")
    add_callout(
        doc,
        "Recommendation",
        "Use two deliberately different models: a normalized PostgreSQL source of truth for integrity, then an ETL-built star schema for dashboard speed and simpler OLAP. Do not force one schema to serve both workloads.",
        "finding",
    )
    add_text(
        doc,
        "The specification can be represented cleanly, but the strongest solution is not just a collection of boxes and foreign keys. It is a chain of traceable decisions: define the grain of each fact, separate stored evidence from derived awards, state ambiguous rules, assign every integrity rule to an enforcement mechanism, and design the dashboard around questions rather than charts. That chain follows the course design sequence from ER modelling to relational design and normalization, then to data warehousing [C1, slides 19–22].",
    )
    add_text(doc, "Five design findings", bold_lead="Five design findings")
    add_bullets(
        doc,
        [
            "Result is an associative entity at rider × stage grain; its composite key prevents duplicate completion records.",
            "Adjusted time should be generated from finish time, bonus, and penalty, not typed independently. GC and Young Rider standings are then derived from trusted stage results.",
            "Points and King of the Mountains cannot be reconstructed from finish time. They need scoring-event evidence or an explicitly documented imported-score feed.",
            "Three headline rules are cross-row or temporal and need deferred triggers or controlled load validation: every team has a rider, locations are used by a stage, and an exited rider has no later result.",
            "The brief is written for one edition, while the Tour is annual. Build the assessed single-edition design first; document a migration path to RaceEdition rather than quietly changing the stated keys.",
        ],
    )
    add_callout(
        doc,
        "Evidence boundary",
        "No rider dataset, BI workbook, or executed benchmark was supplied. The report therefore proposes queries, expected trade-offs, and an evidence protocol, but leaves rankings, screenshots, and response-time cells explicitly pending.",
        "risk",
    )
    add_figure(doc, ASSET_DIR / "design_flow.png", "Figure 1. Traceable design workflow used for the project", width=6.7)


def add_scope_and_requirements(doc: Document) -> None:
    add_heading(doc, "1. Scope, requirements, and design assumptions")
    add_text(
        doc,
        "The baseline scope is one Tour de France edition, matching the project’s declared identifiers: country code, team name, rider bib number, location name, and stage ID [P, pp. 1–3]. The model includes every required entity and adds only the minimum supporting structures needed to make classifications auditable. Proposed extensions are labelled; they are not smuggled in as facts.",
    )
    add_heading(doc, "1.1 Requirement reading method", 2)
    add_numbered(
        doc,
        [
            "Extract nouns as candidate entities and verbs as candidate relationships; then challenge whether a noun is an entity, value, event, or derived result.",
            "For every relationship, write minimum and maximum participation at both ends, using the (min, max) discipline introduced in the ER lecture [C2, slides 16–18].",
            "Identify rules that an ER diagram cannot capture and carry them into a constraint catalogue. This mirrors the tutorial’s instruction to state such rules in English [C5, p. 2].",
            "Record uncertain interpretations in a decision log and obtain owner confirmation before data loading or dashboard calculation.",
        ],
    )

    add_heading(doc, "1.2 Traceability and decisions", 2)
    rows = [
        ["Country", "May exist without teams or riders; ISO alpha-3 identifies it.", "COUNTRY with CHAR(3) PK; no reverse mandatory constraint.", "Specified"],
        ["Team", "Exactly one country; at least one rider.", "FK to COUNTRY; deferred cross-row validation for ≥1 rider.", "Specified"],
        ["Rider", "Exactly one team and one represented country; nationality may differ.", "Two independent NOT NULL FKs; never infer rider country from team.", "Specified"],
        ["Leader / roles", "Background says a designated leader and specialist roles.", "Optional TEAM_LEADER and RIDER_ROLE extension; confirm whether assessed data records them.", "Confirm"],
        ["Location", "Exactly one country; store only if used as a stage endpoint.", "Two role-labelled Stage FKs plus an orphan-location audit/deferred trigger.", "Specified"],
        ["Stage", "One day, one start, one finish, positive distance, four types.", "NOT NULL FKs; domain CHECK; positive distance; start ≠ finish assumption.", "Mostly specified"],
        ["Result", "One row for every stage completer; finish, rank, bonus, penalty.", "Composite PK (stage_id, bib_no); unique (stage_id, rank); generated adjusted time.", "Specified"],
        ["Exit", "Withdrawal before a stage; no results from that stage onward.", "At most one exit per rider; temporal trigger/audit based on stage date/sequence.", "Specified"],
        ["Best Team", "Lowest aggregate time using three fastest team riders.", "Assume three fastest finishers per team per stage, then cumulative; confirm rule.", "Ambiguous"],
        ["Points / KOM", "Required dashboard comparisons; scoring mechanics absent.", "Add scoring events and rider points, or document an imported official-score feed.", "Gap"],
        ["Multiple years", "Background calls the Tour annual; IDs are stated as unique.", "Single 2026 baseline; future RaceEdition migration documented separately.", "Scope choice"],
    ]
    add_table(doc, ["Topic", "Requirement / observation", "Design response", "Status"], rows, widths=[1.15, 2.15, 2.55, 0.8], font_size=7.6)

    add_heading(doc, "1.3 Questions to settle before implementation", 2)
    add_bullets(
        doc,
        [
            "Does ‘three fastest riders’ mean the three best adjusted times in each stage (recommended), or three riders selected over the whole race? What happens when fewer than three riders finish?",
            "Are bonuses and penalties always non-negative seconds, and can they be revised after a stage is closed?",
            "What is the authoritative eligibility cut-off for Young Rider, and is age evaluated once at race start?",
            "Will Points/KOM data arrive as individual scoring events, stage totals, or published standings? Event-level data is most auditable.",
            "Must the database support only 2026, or preserve multiple editions and team/rider history?",
            "Can two historical stages share a date, and is the supplied stage ID guaranteed to represent race order? If not, add stage_sequence.",
        ],
    )
    add_callout(doc, "Decision gate", "Freeze these interpretations in a one-page rulebook before writing calculated fields in SQL or the BI tool. Otherwise two correct-looking dashboards can disagree.", "decision")


def add_conceptual_design(doc: Document) -> None:
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(landscape)
    add_heading(doc, "2. Conceptual design")
    add_text(
        doc,
        "The conceptual model treats keys and participation as business claims, not drawing decoration. It follows the course guidance that entities map real-world objects, relationships associate them, and cardinality/participation encode allowed relationship counts [C2, slides 5–18]. Result and RiderScore are associative entities because their facts exist only for a particular combination of parents; derived values are visibly marked rather than stored as independent truth.",
    )

    add_heading(doc, "2.1 Core ER model", 2)
    add_figure(doc, ASSET_DIR / "core_er.png", "Figure 2. Core ER design with role names and min..max participation", width=9.8)

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(portrait)
    add_text(
        doc,
        "Read each endpoint independently. For example, a country participates in zero or many team relationships, but each team participates in exactly one. A team must participate in one or many rider relationships, while each rider participates in exactly one. Stage uses Location twice under distinct start and finish roles. The combined rule that every Location must be used in at least one of those roles is not expressible as two independent minimums.",
    )
    add_heading(doc, "2.2 Entity and relationship rationale", 2)
    rows = [
        ["COUNTRY", "country_code", "Independent reference entity; optional participation preserves countries with no current team/rider."],
        ["TEAM", "team_name", "Team country is a direct relationship. Team name is accepted as key because the brief declares uniqueness."],
        ["RIDER", "bib_no", "Rider-to-team and rider-to-country are independent mandatory relationships."],
        ["LOCATION", "location_name", "Single country owner; participates in Stage twice with start and finish role names."],
        ["STAGE", "stage_id", "Holds date, type, distance, and endpoint relationships. Add stage_sequence if ID is not ordered."],
        ["RESULT", "stage_id + bib_no", "Association between Rider and Stage with finish facts. (stage_id, rank) is an alternate key if crossing order is unique."],
        ["RIDER_EXIT", "bib_no", "At most one withdrawal per rider. The exit-before Stage FK gives the first stage for which a result is prohibited."],
        ["CLASSIFICATION", "classification_code", "Defines award semantics and calculation direction; does not duplicate a computed winner."],
        ["SCORING_EVENT", "event_id", "Captures where/why Points or KOM points arose; links to one Stage and one Classification."],
        ["RIDER_SCORE", "event_id + bib_no", "Associative evidence for points awarded to a rider at one event."],
    ]
    add_table(doc, ["Entity", "Identifier", "Design reason"], rows, widths=[1.25, 1.45, 4.0], font_size=7.8)

    add_heading(doc, "2.3 Classification and award model", 2)
    add_figure(doc, ASSET_DIR / "classification_model.png", "Figure 3. Classification design: raw evidence feeds derived standings and awards", width=6.7)
    add_text(
        doc,
        "GC, Young Rider, and the proposed Best Team calculation reuse adjusted stage times. Points and KOM use a different measure, so a single generic ‘classification_value’ column would blur incompatible grains and units. Event-level RiderScore records preserve provenance; views aggregate them into standings. Classification metadata controls whether lower time or higher points is better. This separation also prevents a final-winner row from contradicting the underlying facts.",
    )
    add_callout(doc, "Important", "A Classification row defines a rule. An Award is the top-ranked subject produced by that rule at a chosen stage horizon. Treating the award as derived is a design decision, not an omitted entity.", "note")

    add_heading(doc, "2.4 Constraints outside the ER drawing", 2)
    add_table(
        doc,
        ["Rule", "Why ER/FK is insufficient", "Proposed enforcement"],
        [
            ["Every team has ≥1 rider", "Rider→Team FK protects riders, not a newly inserted team with zero riders.", "Load team and roster in one transaction; deferred constraint trigger validates at commit."],
            ["Every location is an endpoint", "A location may have zero starts but one finish, or vice versa; the minimum applies across a union.", "Deferred trigger or periodic orphan audit; restrict direct writes to controlled procedures."],
            ["No result at/after exit stage", "The rule compares rows across Result, Exit, and Stage chronology.", "BEFORE trigger on Result and Exit; repeat as ETL data-quality assertion."],
            ["Every completer has a result", "Absence alone does not distinguish incomplete load, DNS, DNF, or exit.", "Close-stage procedure that validates active roster; optional rider_stage_status extension."],
            ["Exactly one team leader", "At-most-one is easy; at-least-one is a cross-row requirement during loading.", "Optional TEAM_LEADER relation plus deferred completion check."],
        ],
        widths=[1.5, 2.45, 2.75],
        font_size=7.7,
    )


def add_relational_design(doc: Document) -> None:
    add_heading(doc, "3. Relational schema and integrity strategy")
    add_text(
        doc,
        "Entity types become relations; many-to-many or dependent relationships become relations; and one-to-many relationships normally become foreign keys on the many side [C2, slides 22–29]. Primary, foreign-key, domain, CHECK, and trigger constraints are then layered so invalid states fail close to their source [C3, slides 12 and 24–31].",
    )
    add_heading(doc, "3.1 Proposed normalized relations", 2)
    schemas = [
        ["COUNTRY", "country_code CHAR(3)", "country_name", "PK country_code; UNIQUE country_name; ISO-format CHECK"],
        ["TEAM", "team_name TEXT", "country_code", "PK team_name; FK country_code NOT NULL"],
        ["RIDER", "bib_no INT", "rider_name, date_of_birth, team_name, country_code", "PK bib_no; FKs NOT NULL; dob < stage date audit"],
        ["LOCATION", "location_name TEXT", "country_code", "PK location_name; FK country_code NOT NULL"],
        ["STAGE", "stage_id INT", "stage_date, start_location, finish_location, distance_km, stage_type", "PK; FKs; distance > 0; start ≠ finish; stage_type domain"],
        ["STAGE_RESULT", "stage_id + bib_no", "finish_seconds, finish_rank, bonus_seconds, penalty_seconds, adjusted_seconds", "Composite PK; UNIQUE(stage_id, rank); nonnegative checks; adjusted generated"],
        ["RIDER_EXIT", "bib_no", "exit_before_stage_id, reason", "PK/FK bib_no; stage FK; reason nonblank"],
        ["CLASSIFICATION", "classification_code", "name, subject_kind, measure_kind, ranking_direction, eligibility_rule", "PK; unique name; controlled domains"],
        ["SCORING_EVENT", "event_id BIGINT", "stage_id, classification_code, event_label, route_km", "PK; FKs; only POINTS/KOM-compatible classifications"],
        ["RIDER_SCORE", "event_id + bib_no", "points_awarded", "Composite PK; FKs; points ≥ 0"],
    ]
    add_table(doc, ["Relation", "Primary key", "Other attributes", "Core constraints"], schemas, widths=[1.18, 1.45, 2.55, 1.7], font_size=7.25)

    add_heading(doc, "3.2 Constraint catalogue", 2)
    constraints = [
        ["C01", "ISO country code is exactly three uppercase letters.", "CHECK country_code ~ '^[A-Z]{3}$'", "Insert/update"],
        ["C02", "Each team/rider/location refers to an existing country.", "NOT NULL foreign keys, ON DELETE RESTRICT", "Immediate"],
        ["C03", "Each rider belongs to exactly one existing team.", "NOT NULL Rider.team_name FK", "Immediate"],
        ["C04", "Team has at least one rider.", "DEFERRABLE constraint trigger / roster-finalization procedure", "Commit"],
        ["C05", "Stage type is flat, hilly, mountain, or individual time-trial.", "ENUM or CHECK", "Immediate"],
        ["C06", "Stage endpoints exist, differ, and distance is positive.", "FKs + row CHECKs", "Immediate"],
        ["C07", "At most one result per rider-stage and one rank per stage.", "PK(stage_id,bib_no) + UNIQUE(stage_id,finish_rank)", "Immediate"],
        ["C08", "Adjusted = finish − bonus + penalty.", "Generated stored column (or view)", "Automatic"],
        ["C09", "A rider exits at most once.", "RiderExit PK(bib_no)", "Immediate"],
        ["C10", "No result from exit stage onward.", "Temporal trigger using stage_date/stage_sequence", "Write + audit"],
        ["C11", "Every stored location is used.", "Deferred union-of-roles trigger or orphan audit", "Commit/batch"],
        ["C12", "Points are tied to an event, rider, stage, and classification.", "ScoringEvent/RiderScore FKs and domain checks", "Immediate"],
    ]
    add_table(doc, ["ID", "Business rule", "Mechanism", "Timing"], constraints, widths=[0.48, 2.8, 2.55, 0.85], font_size=7.4)

    add_heading(doc, "3.3 Delete, update, and load policy", 2)
    add_text(
        doc,
        "Use RESTRICT for reference/master data (Country, Team, Rider, Location) so an accidental delete cannot erase race history. CASCADE is defensible from Stage to StageResult only if stage deletion is itself tightly restricted; otherwise prefer soft closure/status and explicit repair transactions. Foreign-key delete behavior is part of the logical design, not an afterthought [C3, slides 26–28].",
    )
    add_bullets(
        doc,
        [
            "Load order: Country → Team/Location → Rider → Stage → Result/Exit → scoring events/scores.",
            "Use one transaction for a complete roster so the deferred ‘team has a rider’ constraint is evaluated only at commit.",
            "Close each stage after row counts, rank uniqueness, exit chronology, and arithmetic checks pass; prevent silent edits after closure except through an audited correction workflow.",
            "Keep source file name, ingestion timestamp, and batch ID in staging/audit tables even if they are not part of the assessed ER model.",
        ],
    )
    add_callout(doc, "Design choice", "Do not use NULL Result rows to mean ‘did not finish.’ Presence means completed; absence needs an explicit operational status or a documented exit. This keeps the fact grain unambiguous.", "decision")


def add_normalization(doc: Document) -> None:
    add_heading(doc, "4. Functional dependencies and normalization")
    add_text(
        doc,
        "The normalization argument starts with semantics, not sample rows. A functional dependency holds because the business rule guarantees equal determinants imply equal dependents; one dataset can refute an FD but cannot prove it [C4, slides 10–14]. The proposed decomposition avoids repeated country, team, rider, and location descriptions and therefore the update/insertion/deletion anomalies illustrated in the schema-refinement lecture [C4, slides 3–8].",
    )
    add_heading(doc, "4.1 Main functional dependencies", 2)
    fds = [
        ["COUNTRY", "country_code → country_name", "country_code"],
        ["TEAM", "team_name → country_code", "team_name"],
        ["RIDER", "bib_no → rider_name, date_of_birth, team_name, country_code", "bib_no"],
        ["LOCATION", "location_name → country_code", "location_name"],
        ["STAGE", "stage_id → stage_date, start_location, finish_location, distance_km, stage_type", "stage_id"],
        ["STAGE_RESULT", "(stage_id,bib_no) → finish_seconds, finish_rank, bonus_seconds, penalty_seconds; (stage_id,finish_rank) → bib_no,…", "(stage_id,bib_no); candidate (stage_id,finish_rank)"],
        ["RIDER_EXIT", "bib_no → exit_before_stage_id, reason", "bib_no"],
        ["CLASSIFICATION", "classification_code → name, subject_kind, measure_kind, ranking_direction, eligibility_rule", "classification_code"],
        ["SCORING_EVENT", "event_id → stage_id, classification_code, event_label, route_km", "event_id"],
        ["RIDER_SCORE", "(event_id,bib_no) → points_awarded", "(event_id,bib_no)"],
    ]
    add_table(doc, ["Relation", "Non-trivial FD(s)", "Candidate key(s)"], fds, widths=[1.3, 3.55, 1.85], font_size=7.45)

    add_heading(doc, "4.2 BCNF conclusion", 2)
    add_text(
        doc,
        "Within each listed base relation, every non-trivial determinant is a candidate key; the relations are therefore in BCNF under the stated FDs, and consequently in 3NF. The decomposition is lossless in the practical ER-to-relational sense because parent identifiers are preserved as foreign keys in child relations; joining a child to its parent on that key restores the associated descriptive facts without inventing tuples. The course notes define the central lossless goal as being able to reconstruct the original relation by join [C4, slide 7].",
    )
    add_text(
        doc,
        "The most important avoided transitive dependencies are bib_no → team_name → team_country and stage_id → start_location → start_country. Team country and location country remain in their own relations; a report obtains them by joins. Keeping rider nationality separately is not redundancy because it describes a different relationship and is explicitly allowed to differ from team country.",
    )
    add_table(
        doc,
        ["Tempting shortcut", "Anomaly introduced", "Normalized response"],
        [
            ["Store team_country on every Rider", "A team-country change requires many rider updates; rows may disagree.", "Keep Team.country_code; join when needed."],
            ["Store country_name beside every Location/Team/Rider", "Spelling/name updates can diverge across tables.", "Reference Country by ISO code."],
            ["Store cumulative_gc_time on Rider", "Every result correction requires synchronized rewrites.", "Compute/materialize from StageResult with refresh policy."],
            ["Store final award winner in Classification", "Winner may contradict underlying scores and is unknown mid-race.", "Use ranked views/snapshots with an as-of stage."],
            ["One wide result/classification table", "Mixed seconds/points and changing grain create NULLs and unclear keys.", "Separate StageResult from event-level RiderScore."],
        ],
        widths=[1.6, 2.45, 2.65],
        font_size=7.7,
    )
    add_callout(doc, "Caveat", "Normalization is conditional on the declared rules. If location names are not globally unique, or team names change across editions, the keys must change; a syntactically BCNF table built on a false key is still a bad design.", "risk")


def add_analytics_design(doc: Document) -> None:
    add_heading(doc, "5. Analytics architecture and star schema")
    add_text(
        doc,
        "The operational database and warehouse solve different problems. The course distinguishes an OLTP database with current, frequent transactions from an OLAP warehouse with integrated historical data and complex analytical queries [C1, slides 20–21]. The proposed pipeline keeps PostgreSQL normalized, validates each closed stage, then publishes dimensional rows and immutable fact batches to the dashboard model.",
    )
    add_heading(doc, "5.1 Grain first", 2)
    add_callout(doc, "Fact grain", "FACT_RESULT contains exactly one row for one rider who completed one stage. Every measure and dimension key must describe that same event.", "finding")
    add_figure(doc, ASSET_DIR / "star_schema.png", "Figure 4. Proposed result star with a role-playing Country dimension", width=6.7)
    add_text(
        doc,
        "The star intentionally repeats dimension keys that are reachable through other dimensions. This controlled denormalization makes country, team, rider, and stage filters direct and preserves historical context. Use surrogate keys for slowly changing dimensions and keep the operational natural keys as business-key attributes. For the single-edition prototype, Type 1 dimensions are sufficient; for multi-year history, Rider and Team should use Type 2 validity ranges when affiliations or names change.",
    )
    add_heading(doc, "5.2 Dimensions, measures, and conformance", 2)
    add_table(
        doc,
        ["Object", "Contents", "Design notes"],
        [
            ["FACT_RESULT", "finish, bonus, penalty, adjusted seconds; finish rank; completed_count", "Additive seconds/count across stages; rank is non-additive."],
            ["DIM_RIDER", "bib, name, DOB, age-at-start, young eligibility", "Eligibility must use the confirmed race cut-off, not today’s age."],
            ["DIM_TEAM", "team name, leader/roster descriptors", "Surrogate key supports rebranding/history."],
            ["DIM_STAGE", "stage ID/sequence/date/type/distance/start/finish", "Endpoint names/countries may be flattened for BI convenience."],
            ["DIM_COUNTRY", "ISO code and name", "Role-playing keys distinguish rider nationality and team country."],
            ["FACT_CLASSIFICATION_SCORE", "event × rider points", "Recommended second fact for Points/KOM provenance; conformed Rider/Stage/Classification dimensions."],
        ],
        widths=[1.7, 2.45, 2.55],
        font_size=7.7,
    )

    add_heading(doc, "5.3 ETL quality gates", 2)
    add_numbered(
        doc,
        [
            "Extract closed operational stages by immutable batch ID and capture source counts/checksums.",
            "Validate keys, domains, result arithmetic, unique stage rank, active-roster coverage, and exit chronology; quarantine failures.",
            "Upsert dimensions by natural key, assigning stable surrogate keys and explicit Unknown members for late-arriving references.",
            "Load FACT_RESULT idempotently using the operational (stage_id, bib_no) as the source event key.",
            "Reconcile total row count and sum(adjusted_seconds) between source and fact for each stage, then publish the batch.",
            "Refresh classification views/snapshots and the BI semantic model; display data freshness and the latest closed stage.",
        ],
    )
    add_table(
        doc,
        ["Quality test", "Pass condition", "Failure action"],
        [
            ["Uniqueness", "No duplicate rider-stage or stage-rank records", "Reject batch; investigate source mapping"],
            ["Arithmetic", "adjusted = finish − bonus + penalty for every row", "Reject row/batch; never silently repair"],
            ["Chronology", "No result at or after recorded exit-before stage", "Quarantine rider history; resolve rule/date"],
            ["Completeness", "Expected active riders reconcile to completers/non-completers", "Do not close stage until explained"],
            ["Referential", "Every fact key maps to a dimension or explicit Unknown", "Late-arriving queue; alert owner"],
            ["Reconciliation", "Source/fact counts and adjusted-time totals match", "Rollback publish step"],
        ],
        widths=[1.25, 3.0, 2.45],
        font_size=7.7,
    )


def add_dashboard_design(doc: Document) -> None:
    add_heading(doc, "6. Dashboard and OLAP design")
    add_text(
        doc,
        "The dashboard should answer a small set of decisions: who leads and why, where rankings changed, how classifications differ, which teams accumulate the lowest adjusted time, and how participation differs by country. The interface below is a wireframe, not a fabricated screenshot.",
    )
    add_figure(doc, ASSET_DIR / "dashboard_wireframe.png", "Figure 5. Proposed overview page and interaction hierarchy (illustrative values only)", width=6.7)
    add_heading(doc, "6.1 Page architecture", 2)
    add_table(
        doc,
        ["Page", "Primary question", "Recommended views", "Interaction"],
        [
            ["1. Race overview", "Who leads; how did the race evolve?", "KPIs, cumulative gap lines, stage profile, rank movement", "Stage-range brush; click rider/team; drill-through"],
            ["2. Classifications", "Do the same riders dominate GC, Points, KOM, Young?", "Ranked bars, slope/rank-change chart, stage contribution table", "Classification selector; as-of stage; top-N"],
            ["3. Teams & countries", "Where do teams/riders come from and which teams perform?", "Team time matrix, nationality/team-country bars, optional map", "Cross-filter by team country vs rider nationality"],
            ["4. Rider detail", "What created one rider’s standing?", "Stage-by-stage adjusted time, bonuses/penalties, scores, exit context", "Tooltip → drill-through; return preserving filters"],
        ],
        widths=[1.15, 2.05, 2.25, 1.3],
        font_size=7.55,
    )
    add_heading(doc, "6.2 OLAP behaviour", 2)
    add_bullets(
        doc,
        [
            "Drill down: race total → stage → rider result or scoring event; drill through to the rider detail page.",
            "Roll up: rider results → team → team country, and stage → stage type or date range.",
            "Slice: one classification, stage type, team, rider nationality, or active/exited status.",
            "Dice: compare selected riders/teams across a selected stage range and classification set.",
            "As-of analysis: every classification visual uses a selected/latest closed stage so mid-race and final rankings are not mixed.",
        ],
    )
    add_heading(doc, "6.3 Visual and interaction standards", 2)
    add_bullets(
        doc,
        [
            "Use colour consistently by classification; never rely on colour alone. Add labels, icons, or line styles for accessibility.",
            "Show seconds as hh:mm:ss in labels while retaining numeric seconds in the model for correct aggregation.",
            "Use rank charts with rank 1 at the top; make directionality explicit for points (higher is better) versus time (lower is better).",
            "Keep default views to top-N plus a searchable selector. Dense all-rider charts create false comparison precision.",
            "Expose metric definitions, young-rider eligibility, Best Team rule, latest refresh time, and stage-closure status in a context panel.",
            "Prefer bars/lines/tables for comparison. Use a map only when geography answers a question; country bars remain the accessible fallback.",
        ],
    )
    add_callout(doc, "Acceptance test", "A user should be able to select a stage, identify the GC leader, see the time gap, trace the change to stage results/bonuses/penalties, and return to the overview without losing context.", "finding")


def add_queries_and_evidence(doc: Document) -> None:
    add_heading(doc, "7. Calculation logic and analytical evidence")
    add_text(
        doc,
        "Calculations should be written first as database views and tested on small hand-checkable fixtures. The BI layer may format and filter them, but it should not re-implement core award rules independently. The queries below are representative design artefacts; table/column names can be adjusted to the final implementation.",
    )

    gc_sql = r"""
WITH stage_count AS (
    SELECT COUNT(*) AS n FROM stage
), rider_totals AS (
    SELECT r.bib_no,
           SUM(r.adjusted_seconds) AS cumulative_adjusted_seconds,
           COUNT(*) AS stages_completed
    FROM stage_result r
    GROUP BY r.bib_no
)
SELECT t.bib_no, rd.rider_name, t.cumulative_adjusted_seconds,
       RANK() OVER (ORDER BY t.cumulative_adjusted_seconds, t.bib_no) AS gc_rank
FROM rider_totals t
JOIN rider rd USING (bib_no)
CROSS JOIN stage_count s
WHERE t.stages_completed = s.n;
"""
    add_code_block(doc, gc_sql, "GC at race completion: exclude riders without every stage")
    add_text(
        doc,
        "For an as-of-stage dashboard, replace total stage count with the number of closed stages up to the selected sequence and filter Result to the same horizon. A tie policy must be confirmed; bib number above is only a deterministic display fallback, not a sporting rule.",
    )

    team_sql = r"""
WITH ranked AS (
    SELECT sr.stage_id, r.team_name, sr.bib_no, sr.adjusted_seconds,
           ROW_NUMBER() OVER (
               PARTITION BY sr.stage_id, r.team_name
               ORDER BY sr.adjusted_seconds, sr.finish_rank
           ) AS within_team
    FROM stage_result sr
    JOIN rider r USING (bib_no)
), stage_team AS (
    SELECT stage_id, team_name,
           SUM(adjusted_seconds) FILTER (WHERE within_team <= 3) AS team_stage_seconds,
           COUNT(*) FILTER (WHERE within_team <= 3) AS counted_riders
    FROM ranked
    GROUP BY stage_id, team_name
)
SELECT team_name, SUM(team_stage_seconds) AS cumulative_team_seconds
FROM stage_team
WHERE counted_riders = 3
GROUP BY team_name
ORDER BY cumulative_team_seconds;
"""
    add_code_block(doc, team_sql, "Best Team under the recommended per-stage top-three interpretation")
    add_callout(doc, "Rule to confirm", "The SQL intentionally excludes stage-team rows with fewer than three finishers. The project owner must decide whether that disqualifies the team, uses a fallback rule, or indicates incomplete data.", "risk")

    points_sql = r"""
SELECT c.classification_code, rs.bib_no,
       SUM(rs.points_awarded) AS points,
       DENSE_RANK() OVER (
           PARTITION BY c.classification_code
           ORDER BY SUM(rs.points_awarded) DESC
       ) AS classification_rank
FROM rider_score rs
JOIN scoring_event e USING (event_id)
JOIN classification c USING (classification_code)
GROUP BY c.classification_code, rs.bib_no;
"""
    add_code_block(doc, points_sql, "Points/KOM standings from auditable scoring events")

    add_heading(doc, "7.1 Normalized versus star-schema experiment", 2)
    add_text(
        doc,
        "The final report must compare query complexity and response time. Do not report a single stopwatch run. Use the same PostgreSQL version, hardware, dataset, metric definitions, and equivalent indexes; record plans with EXPLAIN (ANALYZE, BUFFERS). Run each query repeatedly after a controlled warm-up, randomize query order, and report median and p95 wall time with row counts.",
    )
    add_table(
        doc,
        ["Query", "Normalized path", "Star path", "Expected—not measured—effect"],
        [
            ["Cumulative time by rider", "Result → Rider (and Team/Country for filters)", "FactResult → selected dimensions", "Similar aggregation; star removes chained joins"],
            ["Best Team", "Result → Rider → Team, rank top 3 per stage", "FactResult has team_key; same window logic", "Simpler joins; window still dominates"],
            ["Stage performance", "Result → Stage → two Location → Country roles", "FactResult → DimStage", "Star likely much simpler and faster to scan"],
            ["Country participation", "Country ← Team and Country ← Rider branches", "Dimension counts / distinct fact keys", "Star convenient, but avoid double-counting facts"],
        ],
        widths=[1.3, 2.15, 1.75, 1.5],
        font_size=7.45,
    )
    add_heading(doc, "7.2 Benchmark result template", 2)
    add_table(
        doc,
        ["Query", "Rows", "Normalized SQL lines / joins", "Star SQL lines / joins", "Median ms", "p95 ms", "Plan observation"],
        [
            ["GC cumulative", "[run]", "[run]", "[run]", "N: [ ] / S: [ ]", "N: [ ] / S: [ ]", "[EXPLAIN finding]"],
            ["Best Team", "[run]", "[run]", "[run]", "N: [ ] / S: [ ]", "N: [ ] / S: [ ]", "[EXPLAIN finding]"],
            ["Stage analysis", "[run]", "[run]", "[run]", "N: [ ] / S: [ ]", "N: [ ] / S: [ ]", "[EXPLAIN finding]"],
        ],
        widths=[1.0, 0.55, 1.25, 1.15, 0.9, 0.9, 1.15],
        font_size=6.7,
    )
    add_text(
        doc,
        "Interpretation standard: conclude only what the plans and repeated timings support. The likely result is that the normalized schema offers clearer operational integrity and smaller updates, while the star offers simpler BI joins and predictable scans. Dataset size, indexes, caching, and materialization may narrow or reverse a timing difference.",
    )


def add_validation_and_delivery(doc: Document) -> None:
    add_heading(doc, "8. Validation, delivery plan, and risks")
    add_heading(doc, "8.1 Test strategy", 2)
    tests = [
        ["Happy path", "Load 2 countries, 1 team, 3 riders, 2 locations, 1 stage and 3 results.", "All rows accepted; adjusted times and ranks correct."],
        ["Nationality independence", "Team country FRA; rider country DEU.", "Accepted; dashboard distinguishes both country roles."],
        ["Orphan country", "Insert a country with no team/rider.", "Accepted, as explicitly required."],
        ["Empty team", "Commit a Team with no Rider.", "Commit rejected by deferred rule."],
        ["Bad stage", "Negative distance, same start/finish, invalid type.", "Rejected by CHECK/domain constraints."],
        ["Duplicate result", "Same rider-stage twice or two riders same rank.", "Rejected by PK/alternate key."],
        ["Exit chronology", "Exit before stage 5, then insert stage 5/6 result.", "Rejected; earlier results remain allowed."],
        ["Arithmetic", "Finish 10,000; bonus 10; penalty 5.", "Generated adjusted time = 9,995 seconds."],
        ["GC eligibility", "One rider missing a completed stage.", "Excluded from final GC query."],
        ["Points provenance", "Award points without valid event/rider.", "Rejected by FKs."],
        ["ETL idempotence", "Load the same closed-stage batch twice.", "Fact row count and sums unchanged."],
        ["Dashboard drill", "Filter stage/type/team then drill to rider.", "Detail reconciles exactly to source rows and retains context."],
    ]
    add_table(doc, ["Test", "Fixture / action", "Expected result"], tests, widths=[1.2, 3.05, 2.45], font_size=7.6)

    add_heading(doc, "8.2 Work plan aligned to deliverables", 2)
    add_table(
        doc,
        ["Phase", "Output", "Quality gate", "Target"],
        [
            ["Requirements", "Rulebook, assumptions, traceability matrix", "All ambiguities have owner/status", "Before ER freeze"],
            ["Task 1", "ER diagram + cardinality rationale", "Every required entity present; unmodelled rules listed", "21 Sep 2026"],
            ["Task 2", "DDL, FD set, 3NF/BCNF justification, tests", "Fresh database build + negative tests pass", "12 Oct 2026"],
            ["Data / ETL", "Staging, conformed dimensions, star facts", "Stage reconciliation and idempotence pass", "Before dashboard freeze"],
            ["Task 3", "BI workbook + demo script", "All five required dashboard capabilities demonstrated", "11 Nov 2026"],
            ["Task 4", "Final report, screenshots, SQL, benchmark appendix", "Numbers reproducible; citations and AI disclosure complete", "13 Nov 2026"],
        ],
        widths=[1.05, 2.35, 2.55, 0.85],
        font_size=7.55,
    )
    add_text(
        doc,
        "For group work, assign ownership by artefact but require peer review across boundaries: one person owns the rulebook/ER, one DDL and tests, one ETL/star, and one dashboard/benchmark narrative. The owner of a downstream artefact must review its upstream contract; otherwise schema, ETL, and dashboard definitions drift.",
    )

    add_heading(doc, "8.3 Risk register", 2)
    add_table(
        doc,
        ["Risk", "Likelihood / impact", "Early signal", "Mitigation"],
        [
            ["Classification rules underspecified", "High / High", "Different totals in SQL and BI", "Freeze formula, tie, eligibility, and as-of-stage rules"],
            ["Single-edition keys reused later", "Medium / High", "Duplicate bib/stage/team during next load", "Document RaceEdition migration before multi-year ingestion"],
            ["Missing/non-finisher ambiguous", "High / High", "Active roster count ≠ result count without exit", "Stage closure control and optional status table"],
            ["Dashboard double counting", "Medium / High", "Counts change when classification fact joins result fact", "Separate facts; conformed dimensions; measures respect grain"],
            ["Benchmark bias", "Medium / Medium", "One-off timings or non-equivalent indexes", "Repeat, randomize, report median/p95 and plans"],
            ["Late rule correction", "Medium / High", "Manual overwrites after stage close", "Audited revision batches and downstream refresh"],
            ["Academic integrity", "Low / High", "Uncited reused material or unexplained AI output", "Cite sources; keep disclosure; ensure each member can defend work"],
        ],
        widths=[1.45, 1.1, 2.25, 2.0],
        font_size=7.4,
    )


def add_proposed_findings(doc: Document) -> None:
    add_heading(doc, "9. Proposed result and findings")
    add_callout(
        doc,
        "Result",
        "Adopt the normalized schema in Sections 2–4 as the operational source, publish FACT_RESULT plus a separate classification-score fact for analytics, and make all awards reproducible views over those facts. Treat the rulebook and validation suite as deliverables equal in importance to the diagrams.",
        "finding",
    )
    add_text(
        doc,
        "This result is a design finding rather than a race result. It is supported by requirement coverage, dependency analysis, grain consistency, and enforceability. It deliberately does not claim measured dashboard performance or actual 2026 winners, because no race dataset or executed BI artefact was provided.",
    )
    add_table(
        doc,
        ["Finding", "Reason", "What would falsify / revise it"],
        [
            ["BCNF source is suitable", "Each non-trivial FD determinant is a candidate key under declared rules.", "A false uniqueness assumption or newly discovered dependency."],
            ["Adjusted time should be derived", "Its formula is deterministic and stored copies can drift.", "A requirement for officially overridden adjusted values with audit history."],
            ["Scoring facts must be separate", "Points/KOM have different grain and measure from finish-time results.", "An authoritative input contains only final standings and provenance is out of scope."],
            ["Two-schema architecture is preferable", "OLTP integrity and OLAP convenience reward different structures.", "A tiny static dataset where ETL cost exceeds any analytical benefit."],
            ["Star queries should be simpler", "Fact holds direct dimension keys; fewer chained joins.", "Measured plans show no complexity/performance advantage at actual scale."],
        ],
        widths=[1.7, 2.7, 2.4],
        font_size=7.7,
    )
    add_text(
        doc,
        "The defensible final submission should replace every bracketed benchmark cell with executed evidence, insert genuine dashboard screenshots, append the final ER export and full DDL, and update this conclusion if the evidence disagrees with the expected trade-off.",
    )


def add_appendices(doc: Document) -> None:
    add_heading(doc, "Appendix A. PostgreSQL DDL starter")
    add_text(
        doc,
        "This starter covers the declarative core. Deferred constraint-trigger functions and staging/audit tables should be added after the group confirms chronology, team-completeness, and location-lifecycle policies.",
    )
    ddl = r"""
CREATE TYPE stage_type AS ENUM ('flat', 'hilly', 'mountain', 'individual_time_trial');
CREATE TYPE subject_kind AS ENUM ('rider', 'team');
CREATE TYPE measure_kind AS ENUM ('seconds', 'points');
CREATE TYPE rank_direction AS ENUM ('ascending', 'descending');

CREATE TABLE country (
    country_code CHAR(3) PRIMARY KEY
        CHECK (country_code ~ '^[A-Z]{3}$'),
    country_name TEXT NOT NULL UNIQUE
        CHECK (btrim(country_name) <> '')
);

CREATE TABLE team (
    team_name TEXT PRIMARY KEY CHECK (btrim(team_name) <> ''),
    country_code CHAR(3) NOT NULL
        REFERENCES country(country_code) ON DELETE RESTRICT
);

CREATE TABLE rider (
    bib_no INTEGER PRIMARY KEY CHECK (bib_no > 0),
    rider_name TEXT NOT NULL CHECK (btrim(rider_name) <> ''),
    date_of_birth DATE NOT NULL,
    team_name TEXT NOT NULL REFERENCES team(team_name) ON DELETE RESTRICT,
    country_code CHAR(3) NOT NULL REFERENCES country(country_code) ON DELETE RESTRICT
);

CREATE TABLE location (
    location_name TEXT PRIMARY KEY CHECK (btrim(location_name) <> ''),
    country_code CHAR(3) NOT NULL
        REFERENCES country(country_code) ON DELETE RESTRICT
);

CREATE TABLE stage (
    stage_id INTEGER PRIMARY KEY CHECK (stage_id > 0),
    stage_date DATE NOT NULL,
    start_location TEXT NOT NULL REFERENCES location(location_name) ON DELETE RESTRICT,
    finish_location TEXT NOT NULL REFERENCES location(location_name) ON DELETE RESTRICT,
    distance_km NUMERIC(6,2) NOT NULL CHECK (distance_km > 0),
    stage_type stage_type NOT NULL,
    CHECK (start_location <> finish_location)
);

CREATE TABLE stage_result (
    stage_id INTEGER NOT NULL REFERENCES stage(stage_id) ON DELETE RESTRICT,
    bib_no INTEGER NOT NULL REFERENCES rider(bib_no) ON DELETE RESTRICT,
    finish_seconds INTEGER NOT NULL CHECK (finish_seconds > 0),
    finish_rank INTEGER NOT NULL CHECK (finish_rank > 0),
    bonus_seconds INTEGER NOT NULL DEFAULT 0 CHECK (bonus_seconds >= 0),
    penalty_seconds INTEGER NOT NULL DEFAULT 0 CHECK (penalty_seconds >= 0),
    adjusted_seconds INTEGER GENERATED ALWAYS AS
        (finish_seconds - bonus_seconds + penalty_seconds) STORED,
    PRIMARY KEY (stage_id, bib_no),
    UNIQUE (stage_id, finish_rank),
    CHECK (finish_seconds - bonus_seconds + penalty_seconds >= 0)
);

CREATE TABLE rider_exit (
    bib_no INTEGER PRIMARY KEY REFERENCES rider(bib_no) ON DELETE RESTRICT,
    exit_before_stage_id INTEGER NOT NULL REFERENCES stage(stage_id) ON DELETE RESTRICT,
    reason TEXT NOT NULL CHECK (btrim(reason) <> '')
);

CREATE TABLE classification (
    classification_code VARCHAR(16) PRIMARY KEY,
    classification_name TEXT NOT NULL UNIQUE,
    subject_kind subject_kind NOT NULL,
    measure_kind measure_kind NOT NULL,
    ranking_direction rank_direction NOT NULL,
    eligibility_rule TEXT
);

CREATE TABLE scoring_event (
    event_id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    stage_id INTEGER NOT NULL REFERENCES stage(stage_id) ON DELETE RESTRICT,
    classification_code VARCHAR(16) NOT NULL
        REFERENCES classification(classification_code) ON DELETE RESTRICT,
    event_label TEXT NOT NULL CHECK (btrim(event_label) <> ''),
    route_km NUMERIC(6,2) CHECK (route_km >= 0)
);

CREATE TABLE rider_score (
    event_id BIGINT NOT NULL REFERENCES scoring_event(event_id) ON DELETE RESTRICT,
    bib_no INTEGER NOT NULL REFERENCES rider(bib_no) ON DELETE RESTRICT,
    points_awarded NUMERIC(8,2) NOT NULL CHECK (points_awarded >= 0),
    PRIMARY KEY (event_id, bib_no)
);

CREATE INDEX ix_result_bib_stage ON stage_result (bib_no, stage_id);
CREATE INDEX ix_rider_team ON rider (team_name);
CREATE INDEX ix_score_rider ON rider_score (bib_no, event_id);
"""
    add_code_block(doc, ddl)

    add_heading(doc, "Appendix B. Multi-edition migration path")
    add_text(
        doc,
        "Do not retrofit multiple years by assuming bib number or stage ID remains globally unique. Introduce RaceEdition(race_id, year, start_date, end_date), stable RiderPerson and TeamOrganisation identities, and an Entry/Roster association linking a person and team to one edition and bib. Stage becomes keyed by (race_id, stage_no) or a surrogate stage_key; Result references the edition-specific entry and stage. Location names should gain a surrogate location_id if duplicates across countries or languages are possible. This extension changes keys and should be designed before ingesting a second year.",
    )
    add_table(
        doc,
        ["Single-edition object", "Multi-edition replacement", "Reason"],
        [
            ["Rider(bib_no)", "RiderPerson + RaceEntry(race_id, bib_no, rider_id, team_id)", "Bib numbers are event assignments, not stable identities."],
            ["Team(team_name)", "TeamOrganisation(team_id) + edition entry/name history", "Teams can rebrand or change country."],
            ["Stage(stage_id)", "Stage(race_id, stage_no) or surrogate stage_key", "Stage numbering repeats each year."],
            ["Result(stage_id,bib_no)", "Result(stage_key, entry_key)", "Fact keys must include edition context."],
        ],
        widths=[1.55, 3.4, 1.85],
        font_size=7.7,
    )

    add_heading(doc, "Appendix C. Course-material source map")
    add_text(
        doc,
        "References below point to the supplied course PDFs. They are concise pointers rather than reproduced lecture content, consistent with the course notice not to redistribute materials.",
    )
    sources = [
        ["[P]", "BT5110 Project: Tour de France Database and Analytics System, AY26/27 Semester 1", "pp. 1–4", "Application rules, tasks, deliverables, due dates"],
        ["[C1]", "Lecture 1: Introduction", "slides 6 and 19–22", "Design sequence; database vs warehouse; AI responsibility"],
        ["[C2]", "Lecture 3: Entity-Relationship Model", "slides 5–18; 22–31", "Entities, relationships, keys, cardinality, participation, translation"],
        ["[C3]", "Lecture 2: Relational Model", "slides 12; 24–31", "Integrity, PK/FK/referential integrity, delete policy, CHECK constraints"],
        ["[C4]", "Lecture 4: Schema Refinement—Functional Dependencies", "slides 3–14; 43–47", "Anomalies, lossless decomposition, FD semantics, minimal cover"],
        ["[C5]", "Tutorial 2: Entity-Relationship Modelling", "pp. 1–2", "Requirement-to-ER method and listing constraints outside ER"],
        ["[C6]", "Tutorials 3 and 4", "p. 1 each", "Closure/minimal cover and 3NF/BCNF decomposition workflow"],
    ]
    add_table(doc, ["Ref", "Source", "Location", "Used for"], sources, widths=[0.5, 3.2, 1.25, 1.85], font_size=7.55)

    doc.add_page_break()
    add_heading(doc, "Appendix D. Submission checklist and AI acknowledgement")
    add_table(
        doc,
        ["Check", "Owner", "Status"],
        [
            ["Replace title-page student/group placeholders", "[name]", "□"],
            ["Confirm Best Team, Young Rider, tie, and stage-order rules", "[name]", "□"],
            ["Export final ER and verify cardinalities against rulebook", "[name]", "□"],
            ["Execute DDL from an empty database and retain test evidence", "[name]", "□"],
            ["Replace dashboard wireframe with genuine screenshots", "[name]", "□"],
            ["Run normalized/star benchmark and complete result table", "[name]", "□"],
            ["Refresh Word table of contents and inspect page breaks", "[name]", "□"],
            ["Each group member reviews and can explain submitted material", "all", "□"],
        ],
        widths=[4.7, 1.2, 0.7],
        font_size=7.8,
    )
    add_callout(
        doc,
        "AI disclosure template",
        "Generative AI was used to help structure a design approach, draft editable prose/SQL, and generate code-based diagrams. The group reviewed the project and course sources, verified and revised all design decisions, executed the final database/dashboard tests, and remains responsible for the submitted work. Adapt this wording to the module’s required disclosure format.",
        "note",
    )
    add_text(
        doc,
        "The course explicitly requires AI use to be acknowledged and warns that students remain responsible for understanding and verifying generated content [C1, slide 6]. Keep this acknowledgement only if it accurately describes the group’s actual process.",
    )
    add_heading(doc, "Final group review record", 2)
    add_table(
        doc,
        ["Reviewer", "Scope reviewed", "Date", "Outcome / change made"],
        [
            ["[member 1]", "Requirements + ER", "[date]", "[record]"],
            ["[member 2]", "DDL + normalization", "[date]", "[record]"],
            ["[member 3]", "ETL + star schema", "[date]", "[record]"],
            ["[member 4]", "Dashboard + benchmarks", "[date]", "[record]"],
        ],
        widths=[1.15, 2.0, 1.0, 2.45],
        font_size=7.7,
    )
    add_callout(
        doc,
        "Editing note",
        "All narrative, tables, SQL, labels, and placeholders in the DOCX are editable. The five embedded diagrams can be regenerated after editing the accompanying Python script.",
        "decision",
    )


def build_document() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_toc(doc)
    add_executive_summary(doc)
    add_scope_and_requirements(doc)
    add_conceptual_design(doc)
    add_relational_design(doc)
    add_normalization(doc)
    add_analytics_design(doc)
    add_dashboard_design(doc)
    add_queries_and_evidence(doc)
    add_validation_and_delivery(doc)
    add_proposed_findings(doc)
    add_appendices(doc)

    core = doc.core_properties
    core.title = "Tour de France Database & Analytics System — Design Blueprint"
    core.subject = "BT5110 Data Management and Warehousing project design"
    core.author = "[Student group — update before submission]"
    core.keywords = "ER model, BCNF, PostgreSQL, star schema, OLAP, dashboard, Tour de France"
    core.comments = "Generated by build_tour_de_france_design_report.py from the supplied project brief and course materials."
    doc.save(OUTPUT_FILE)


def validate_docx() -> dict[str, int]:
    if not OUTPUT_FILE.exists() or OUTPUT_FILE.stat().st_size < 250_000:
        raise AssertionError("DOCX is missing or unexpectedly small")
    with ZipFile(OUTPUT_FILE) as archive:
        names = set(archive.namelist())
        required = {"word/document.xml", "[Content_Types].xml", "docProps/core.xml"}
        if not required.issubset(names):
            raise AssertionError("DOCX ZIP package is incomplete")
        xml = archive.read("word/document.xml").decode("utf-8")
        image_count = sum(name.startswith("word/media/") for name in names)
        if image_count < 6:
            raise AssertionError(f"Expected at least six embedded images, found {image_count}")
        for expected in ("Executive design position", "Core ER model", "BCNF conclusion", "Proposed result and findings"):
            if expected not in xml:
                raise AssertionError(f"Expected section missing: {expected}")
        if "[run]" not in xml:
            raise AssertionError("Benchmark evidence placeholders were unexpectedly removed")
    check = Document(OUTPUT_FILE)
    words = sum(len(p.text.split()) for p in check.paragraphs)
    for table in check.tables:
        for row in table.rows:
            for cell in row.cells:
                words += sum(len(p.text.split()) for p in cell.paragraphs)
    if words < 5000:
        raise AssertionError(f"Document is unexpectedly short: {words} words")
    return {"words": words, "tables": len(check.tables), "images": len(check.inline_shapes)}


def render_pdf_for_review() -> Path | None:
    """Render a PDF sibling for visual QA when LibreOffice is available.

    The PDF is a review artefact, not the requested deliverable.  It is retained because
    it makes layout checking easy and does not reduce DOCX editability.
    """
    libreoffice = shutil.which("libreoffice")
    if not libreoffice:
        return None
    review_dir = ROOT / "rendered_review"
    review_dir.mkdir(exist_ok=True)
    result = subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(review_dir), str(OUTPUT_FILE)],
        cwd=ROOT,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=120,
        check=False,
    )
    pdf_path = review_dir / (OUTPUT_FILE.stem + ".pdf")
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError("LibreOffice review render failed:\n" + result.stdout)
    return pdf_path


def main() -> None:
    validate_sources()
    generate_assets()
    build_document()
    stats = validate_docx()
    review_pdf = render_pdf_for_review()
    print(f"Created: {OUTPUT_FILE}")
    print(f"Document QA: {stats['words']:,} words; {stats['tables']} tables; {stats['images']} inline images")
    if review_pdf:
        print(f"Review PDF: {review_pdf}")


if __name__ == "__main__":
    main()
